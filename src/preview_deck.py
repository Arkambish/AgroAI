"""Approximate PNG preview of a .pptx (shape boxes + text) for layout QA.

Not a faithful renderer — it draws each shape's fill/outline and its text with
Pillow so overlaps, collisions and whitespace balance can be inspected visually.

Run: python src/preview_deck.py outputs/deck/AgroAI_Final_Presentation_15min.pptx
"""

import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.text import PP_ALIGN

EMU = 914400.0
DPI = 110

FONT_DIRS = ["/System/Library/Fonts/Supplemental/", "/System/Library/Fonts/",
             "/Library/Fonts/"]
CANDIDATES = {"reg": ["Arial.ttf", "Helvetica.ttc"],
              "bold": ["Arial Bold.ttf", "Arial.ttf", "Helvetica.ttc"]}
_cache = {}


def font(sz, bold=False):
    key = (round(sz), bold)
    if key in _cache:
        return _cache[key]
    for d in FONT_DIRS:
        for name in CANDIDATES["bold" if bold else "reg"]:
            p = Path(d) / name
            if p.exists():
                try:
                    f = ImageFont.truetype(str(p), max(6, int(sz)))
                    _cache[key] = f
                    return f
                except Exception:
                    pass
    f = ImageFont.load_default()
    _cache[key] = f
    return f


def rgb(c, default=(140, 140, 140)):
    try:
        return (c[0], c[1], c[2])
    except Exception:
        return default


def fill_of(sh):
    try:
        ft = str(sh.fill.type)
        if "SOLID" in ft:
            return rgb(sh.fill.fore_color.rgb)
        if "GRADIENT" in ft:
            return rgb(sh.fill.gradient_stops[0].color.rgb)
    except Exception:
        pass
    return None


def line_of(sh):
    try:
        if "SOLID" in str(sh.line.fill.type):
            return rgb(sh.line.color.rgb)
    except Exception:
        pass
    return None


def wrap(draw, text, f, maxw):
    words, lines, cur = text.split(), [], ""
    for w in words:
        t = (cur + " " + w).strip()
        if draw.textlength(t, font=f) <= maxw or not cur:
            cur = t
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def render(pptx_path, outdir):
    prs = Presentation(pptx_path)
    W = int(prs.slide_width / EMU * DPI)
    H = int(prs.slide_height / EMU * DPI)
    outdir.mkdir(parents=True, exist_ok=True)
    paths = []
    for i, slide in enumerate(prs.slides, 1):
        img = Image.new("RGB", (W, H), (255, 255, 255))
        d = ImageDraw.Draw(img)
        for sh in slide.shapes:
            if sh.left is None:
                continue
            x0 = sh.left / EMU * DPI
            y0 = sh.top / EMU * DPI
            x1 = x0 + (sh.width or 0) / EMU * DPI
            y1 = y0 + (sh.height or 0) / EMU * DPI
            fc, lc = fill_of(sh), line_of(sh)
            if fc or lc:
                d.rounded_rectangle([x0, y0, x1, y1], radius=4, fill=fc,
                                    outline=lc, width=1)
            if not sh.has_text_frame:
                continue
            tf = sh.text_frame
            cy = y0 + 3
            for p in tf.paragraphs:
                txt = "".join(r.text for r in p.runs)
                if not txt.strip():
                    cy += 5
                    continue
                r0 = p.runs[0]
                sz = (r0.font.size.pt if r0.font.size else 11) * DPI / 72.0
                fnt = font(sz, bool(r0.font.bold))
                col = rgb(r0.font.color.rgb, (30, 30, 30)) if r0.font.color and \
                    r0.font.color.type is not None else (30, 30, 30)
                maxw = max(20, (x1 - x0) - 8)
                for ln in wrap(d, txt, fnt, maxw):
                    lw = d.textlength(ln, font=fnt)
                    if p.alignment == PP_ALIGN.CENTER:
                        lx = x0 + ((x1 - x0) - lw) / 2
                    elif p.alignment == PP_ALIGN.RIGHT:
                        lx = x1 - lw - 4
                    else:
                        lx = x0 + 4
                    d.text((lx, cy), ln, font=fnt, fill=col)
                    cy += sz * 1.18
                cy += (p.space_after.pt if p.space_after else 2) * DPI / 72.0
        p = outdir / f"slide_{i:02d}.png"
        img.save(p)
        paths.append(p)
    return paths


if __name__ == "__main__":
    src = Path(sys.argv[1])
    out = Path(sys.argv[2] if len(sys.argv) > 2 else "outputs/deck/preview")
    for p in render(src, out):
        print(p)
